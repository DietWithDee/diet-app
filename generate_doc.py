import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE

def create_article():
    doc = docx.Document()

    # Set page margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Styles
    styles = doc.styles

    # Normal Style
    normal_style = styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(11.5)
    normal_style.font.color.rgb = RGBColor(0x33, 0x33, 0x33) # Charcoal
    normal_style.paragraph_format.line_spacing = 1.25
    normal_style.paragraph_format.space_after = Pt(8)

    # Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run_title = p_title.add_run("Let’s Talk About Sugar: Mindful Drinking & Portion Control")
    run_title.font.name = 'Arial'
    run_title.font.size = Pt(22)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(0x1B, 0x4D, 0x3E) # Deep Forest Green
    p_title.paragraph_format.space_after = Pt(4)

    # Subtitle / Metadata
    p_sub = doc.add_paragraph()
    run_sub = p_sub.add_run("Inspired by Diet with Dee  |  Nutrition & Wellness Insights")
    run_sub.font.name = 'Calibri'
    run_sub.font.size = Pt(10.5)
    run_sub.font.italic = True
    run_sub.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
    p_sub.paragraph_format.space_after = Pt(20)

    # Divider line helper
    def add_divider():
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(16)
        run = p.add_run("—" * 35)
        run.font.color.rgb = RGBColor(0xD0, 0xD0, 0xD0)

    # Section Heading Helper
    def add_heading(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(6)
        run = p.add_run(text)
        run.font.name = 'Arial'
        run.font.size = Pt(14)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0x2E, 0x7D, 0x32) # Vibrant Olive/Green
        return p

    # Callout Box Helper
    def add_callout(text, bold_prefix=""):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.4)
        p.paragraph_format.right_indent = Inches(0.4)
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(12)
        if bold_prefix:
            r_b = p.add_run(bold_prefix + " ")
            r_b.font.bold = True
            r_b.font.color.rgb = RGBColor(0x1B, 0x4D, 0x3E)
        r_t = p.add_run(text)
        r_t.font.italic = True
        r_t.font.color.rgb = RGBColor(0x44, 0x44, 0x44)

    # Article Body Content

    # Introduction
    p = doc.add_paragraph()
    p.add_run("When starting a health journey or trying to manage your weight, one of the very first changes most people make is cutting back on added sugar. You stop putting table sugar in your morning brew, skip dessert, and feel great about your dedication. But then weeks go by, and you ask yourself: ")
    r_quote = p.add_run("“Why is my progress stalling?”")
    r_quote.font.bold = True

    p2 = doc.add_paragraph()
    p2.add_run("The secret often lies in our daily routines and liquid intake. Without realizing it, many of us enjoy multiple sweet drinks throughout a busy day on autopilot. While you may not be eating sugar in your main meals, ")
    p2.add_run("your liquid intake might be working overtime!").font.bold = True

    add_divider()

    # Heading 1
    add_heading("Auditing Your Daily History: Frequency Matters")
    
    p = doc.add_paragraph()
    p.add_run("Think back through your history over the past 24 hours. Did you have a bottled drink with lunch, another during a mid-afternoon energy slump, and maybe a third while relaxing in the evening? ")

    add_callout(
        "Sipping 2 or 3 bottled beverages throughout the day can add significant liquid calories and sugar to your daily total without you even noticing. Because we drink them quickly, they don't fill us up the way solid food does.",
        "A Quick Daily Audit:"
    )

    p = doc.add_paragraph()
    p.add_run("It’s easy to forget that liquid calories count just as much as food calories. When you take small sips throughout the day, those numbers quietly compound, which can explain why your weight loss or blood sugar goals feel harder to reach.")

    # Heading 2
    add_heading("It’s About Moderation, Not Deprivation")

    p = doc.add_paragraph()
    p.add_run("Here is the good news: ")
    p.add_run("you don't have to ban drinks entirely from your life.").font.bold = True
    p.add_run(" Nutrition is about balance, sustainability, and mindfulness—not extreme restriction.")

    p = doc.add_paragraph()
    p.add_run("The key is simply being aware of your portion sizes and total daily intake. Enjoying your favorite beverage is completely fine, but being mindful of how often you reach for a bottle helps you stay in control of your health goals.")

    add_callout(
        "“I’m not saying don’t enjoy your drinks—you can! Just remember that drinks count toward your daily sugar and calorie totals. Being mindful of your portion frequency makes all the difference.”",
        "Dietitian's Take:"
    )

    add_divider()

    # Heading 3: Diabetes Plan
    add_heading("Take Control: Blood Sugar Balance Guide")

    p = doc.add_paragraph()
    p.add_run("If you want to take the guesswork out of your daily nutrition, balance your blood sugar, and follow a clear, practical meal plan, check out our guide:")

    # Plain text URL paragraph
    p_url = doc.add_paragraph()
    p_url.paragraph_format.left_indent = Inches(0.3)
    p_url.paragraph_format.space_before = Pt(6)
    p_url.paragraph_format.space_after = Pt(10)
    
    r_label = p_url.add_run("Blood Sugar Balance (Type 2 Diabetes-Friendly Guide): ")
    r_label.font.bold = True
    
    r_link = p_url.add_run("https://paystack.com/buy/blood-sugar-balance-plan")
    r_link.font.name = 'Consolas'
    r_link.font.size = Pt(10.5)
    r_link.font.color.rgb = RGBColor(0x10, 0x4E, 0x8B)

    p_note = doc.add_paragraph()
    p_note.add_run("Copy and paste the link above into your web browser to get your copy of the plan today.")
    p_note.runs[0].font.italic = True
    p_note.runs[0].font.size = Pt(10)
    p_note.runs[0].font.color.rgb = RGBColor(0x66, 0x66, 0x66)

    # Heading 4: Teaser
    add_heading("Cooking Something Special: App Partnership Teaser! 🤫📱")

    p = doc.add_paragraph()
    p.add_run("We know that keeping track of daily portions, liquid calories, and blood sugar balance on your own can take a lot of effort.")

    p_teaser = doc.add_paragraph()
    p_teaser.add_run("That is why we are super excited to tease that ")
    p_teaser.add_run("we are currently cooking something special! ").font.bold = True
    p_teaser.add_run("We have partnered up to bring you a brand-new, dedicated app built to make tracking habits, managing sugar intake, and staying on top of your wellness goals completely seamless.")

    p = doc.add_paragraph()
    p.add_run("Stay tuned—more details and an exclusive sneak peek will be dropping very soon!")

    add_divider()

    # Sign off
    p_sign = doc.add_paragraph()
    p_sign.paragraph_format.space_before = Pt(12)
    run_sign = p_sign.add_run("Have a great weekend!")
    run_sign.font.name = 'Arial'
    run_sign.font.size = Pt(14)
    run_sign.font.bold = True
    run_sign.font.color.rgb = RGBColor(0x1B, 0x4D, 0x3E)

    doc.save("Lets_Talk_About_Sugar.docx")
    print("Successfully generated revised Lets_Talk_About_Sugar.docx")

if __name__ == "__main__":
    create_article()
