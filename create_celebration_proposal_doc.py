import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    """Sets background color for a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Sets inner margins for a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def create_proposal_doc(output_path="DietWithDee_TikTok_Celebration_Ideas_Proposal.docx"):
    doc = docx.Document()

    # Set margins
    for section in doc.sections:
        section.top_margin = Inches(0.9)
        section.bottom_margin = Inches(0.9)
        section.left_margin = Inches(0.9)
        section.right_margin = Inches(0.9)

    # Base styles
    styles = doc.styles
    normal_style = styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = RGBColor(0x2D, 0x37, 0x48) # Slate charcoal
    normal_style.paragraph_format.line_spacing = 1.2
    normal_style.paragraph_format.space_after = Pt(6)

    # Palette
    COLOR_PRIMARY = RGBColor(0x1B, 0x4D, 0x3E)   # Deep Forest Green
    COLOR_SECONDARY = RGBColor(0x2E, 0x7D, 0x32) # Vibrant Emerald Green
    COLOR_ACCENT = RGBColor(0xD9, 0x77, 0x06)    # Amber Gold
    COLOR_MUTED = RGBColor(0x64, 0x74, 0x8B)     # Muted Slate
    COLOR_DARK = RGBColor(0x1E, 0x29, 0x3B)      # Dark Charcoal

    # Helper Functions
    def add_main_title(title_text, subtitle_text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(title_text)
        run.font.name = 'Arial'
        run.font.size = Pt(22)
        run.font.bold = True
        run.font.color.rgb = COLOR_PRIMARY

        p_sub = doc.add_paragraph()
        p_sub.paragraph_format.space_after = Pt(16)
        run_sub = p_sub.add_run(subtitle_text)
        run_sub.font.name = 'Calibri'
        run_sub.font.size = Pt(11.5)
        run_sub.font.italic = True
        run_sub.font.color.rgb = COLOR_MUTED

    def add_h1(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(16)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = 'Arial'
        run.font.size = Pt(15)
        run.font.bold = True
        run.font.color.rgb = COLOR_PRIMARY
        return p

    def add_h2(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = 'Arial'
        run.font.size = Pt(12.5)
        run.font.bold = True
        run.font.color.rgb = COLOR_SECONDARY
        return p

    def add_divider():
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(12)
        run = p.add_run("―" * 45)
        run.font.color.rgb = RGBColor(0xEB, 0xED, 0xF0)

    def add_callout(title, body, bg_hex="F0FDF4", border_hex="2E7D32"):
        table = doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = table.cell(0, 0)
        set_cell_background(cell, bg_hex)
        set_cell_margins(cell, top=120, bottom=120, left=180, right=180)
        
        # Left border styling
        tcPr = cell._tc.get_or_add_tcPr()
        tcBorders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:left w:val="single" w:sz="24" w:space="0" w:color="{border_hex}"/><w:top w:val="none"/><w:right w:val="none"/><w:bottom w:val="none"/></w:tcBorders>')
        tcPr.append(tcBorders)

        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(3)
        r_title = p.add_run(title)
        r_title.font.bold = True
        r_title.font.size = Pt(11)
        r_title.font.color.rgb = COLOR_PRIMARY

        p2 = cell.add_paragraph()
        p2.paragraph_format.space_after = Pt(0)
        r_body = p2.add_run(body)
        r_body.font.italic = True
        r_body.font.size = Pt(10.5)
        r_body.font.color.rgb = COLOR_DARK

        p_space = doc.add_paragraph()
        p_space.paragraph_format.space_before = Pt(4)
        p_space.paragraph_format.space_after = Pt(4)

    # --- DOCUMENT GENERATION CONTENT ---

    add_main_title(
        "TikTok Milestone Celebration Campaign Proposals",
        "Strategy & Campaign Ideas to Drive Follower Growth, Website Traffic & Consultation Sales | Prepared for Dee"
    )

    add_callout(
        "🎯 Strategic Objective",
        "Celebrate our 5,000 / 10,000 TikTok follower milestone by executing an interactive campaign that fulfills three key goals: (1) Exponentially grow TikTok followers, (2) Increase dietwithdee.org website engagement, and (3) Convert active followers into paying meal plan buyers and 1-on-1 consultation clients."
    )

    add_h1("1. Executive Overview & Campaign Philosophy")
    p = doc.add_paragraph()
    p.add_run("Reaching 5,000 and 10,000 TikTok followers is a tremendous milestone for DietWithDee! Rather than doing a standard giveaway that only yields passive likes, our campaign strategy turns this milestone into a ")
    p.add_run("high-converting growth loop").bold = True
    p.add_run(". Every activity is intentionally structured to direct social traffic onto ")
    p.add_run("dietwithdee.org").bold = True
    p.add_run(" where followers discover our consultation services, meal plans, and resources.")

    add_divider()

    add_h1("2. Refined Core Ideas (Optimized for High Conversion & Low Friction)")

    add_h2("Idea A: The 'Find the Diamond' Web Treasure Hunt (Interactive & High Intent)")
    p = doc.add_paragraph()
    p.add_run("• Concept: ").bold = True
    p.add_run("Hide 3 small interactive 'Diamond/Gem' icons across high-value pages on dietwithdee.org (e.g., hidden on the 1-on-1 Consultation page, Meal Plan Catalog, and About Dee page).\n")
    p.add_run("• How it Works: ").bold = True
    p.add_run("Followers visit the website to search for the gems. Clicking a gem opens an instant pop-up revealing a Secret Celebration Code (e.g., ")
    p.add_run("DEEDIAMOND10K").bold = True
    p.add_run(") granting an exclusive 15% discount on consultations/plans OR entry into a mega giveaway.\n")
    p.add_run("• The Viral TikTok Loop: ").bold = True
    p.add_run("To enter the giveaway grand prize (a 1-Month Free Custom Meal Plan + Consultation), followers must screenshot the found diamond page and post it on TikTok/IG tagging @DietWithDee or comment on our pinned video.\n")
    p.add_run("• Why Dee Will Love It: ").bold = True
    p.add_run("100% of participants are forced to browse our sales and service pages. They get exposed to all our paid offerings while having fun!")

    add_h2("Idea B: Frictionless 1-Tap 'Share to Unlock' (Smart Referral Alternative)")
    p = doc.add_paragraph()
    p.add_run("• Concept: ").bold = True
    p.add_run("Traditional referral systems (requiring users to register, create accounts, and track unique links) often suffer from high user drop-off. Instead, we introduce a ")
    p.add_run("1-Tap Social Unlock").bold = True
    p.add_run(" page at dietwithdee.org/celebrate.\n")
    p.add_run("• How it Works: ").bold = True
    p.add_run("Followers land on the page to claim 'Dee's 5-Day Flat-Belly Recipe Teaser PDF'. Clicking a single button ('Share to WhatsApp' or 'Share to TikTok') instantly unlocks the downloadable guide and presents a prompt to book a consultation.\n")
    p.add_run("• Why Dee Will Love It: ").bold = True
    p.add_run("Zero friction for the user, rapid viral sharing across WhatsApp groups and TikTok, and zero complex backend development needed.")

    add_divider()

    add_h1("3. Wild, Out-of-the-Box & Crazy Ideas (High Impact & Viral Potential)")

    add_h2("Idea C: 'Dee's Mystery Meal Plan Wheel' (Spin-to-Win Lead Capture)")
    p = doc.add_paragraph()
    p.add_run("• Concept: ").bold = True
    p.add_run("An interactive digital prize wheel embedded on dietwithdee.org/spin.\n")
    p.add_run("• Prizes: ").bold = True
    p.add_run("1x Grand Prize Free Consultation, 5x 50% Off Meal Plans, 10x Free Grocery Shopping Guides, and 100x Secret Discount Codes.\n")
    p.add_run("• The Lead Hook: ").bold = True
    p.add_run("To spin the wheel, visitors enter their Email & TikTok handle. This builds a targeted email subscriber list of people interested in diet & health, which can be retargeted for consultation bookings.\n")

    add_h2("Idea D: 'Rate My Plate' / The TikTok Meal Roast & Consultation Fix")
    p = doc.add_paragraph()
    p.add_run("• Concept: ").bold = True
    p.add_run("A fun submission form at dietwithdee.org/roast where followers upload a photo of their daily lunch or dinner plate.\n")
    p.add_run("• The TikTok Content: ").bold = True
    p.add_run("Dee selects 10 submitted plates and posts a multi-part TikTok series roasting/evaluating them with expert nutrition advice.\n")
    p.add_run("• The Sales Pitch: ").bold = True
    p.add_run("At the end of every video, Dee says: 'Want a plate built specifically for your body and weight goals? Book your 1-on-1 consultation at dietwithdee.org!'")

    add_h2("Idea E: The Golden Ticket Meal Plan (Charlie & the Chocolate Factory Style)")
    p = doc.add_paragraph()
    p.add_run("• Concept: ").bold = True
    p.add_run("For a 72-hour celebration window, every digital Meal Plan or Consultation guide purchased on dietwithdee.org has a chance of containing a digital 'Golden Ticket' hidden inside the PDF.\n")
    p.add_run("• Golden Ticket Prize: ").bold = True
    p.add_run("3 lucky buyers win a VIP 30-Minute Follow-up Coaching Session with Dee.\n")
    p.add_run("• Impact: ").bold = True
    p.add_run("Creates immense buying urgency and FOMO, transforming passive followers into immediate paying customers.")

    add_h2("Idea F: 10k TikTok Live Consultation Blitz (Flash Sale Event)")
    p = doc.add_paragraph()
    p.add_run("• Concept: ").bold = True
    p.add_run("Dee hosts a celebratory 1-Hour TikTok LIVE stream (e.g. 'Live Q&A: 5 Myths Sabotaging Your Diet Goals').\n")
    p.add_run("• The Flash Sale: ").bold = True
    p.add_run("During the live stream, Dee drops 10 exclusive 'Celebration Consultation Slots' at a 40% discount, available ONLY on dietwithdee.org/live during the stream duration.\n")
    p.add_run("• Impact: ").bold = True
    p.add_run("Real-time trust, immediate urgency, and guaranteed consultation bookings in under 60 minutes.")

    add_h2("Idea G: Community Goal Unlock Bar (Global Gamification)")
    p = doc.add_paragraph()
    p.add_run("• Concept: ").bold = True
    p.add_run("A live progress bar at the top of dietwithdee.org showing total website visits or social shares. When the bar reaches 2,000 total shares, a 'Secret 7-Day Healthy Snack Guide' unlocks for everyone for free for 48 hours.")

    add_divider()

    add_h1("4. Idea Evaluation & Strategy Matrix")
    p = doc.add_paragraph()
    p.add_run("To help Dee choose the best mix of ideas, here is a breakdown by objective, effort level, and revenue potential:")

    # Table creation
    table = doc.add_table(rows=8, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    col_widths = [Inches(2.2), Inches(1.4), Inches(1.2), Inches(1.7)]
    headers = ["Campaign Idea", "Primary Goal", "Tech Effort", "Revenue / Sales Impact"]

    # Header Row
    hdr_cells = table.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].width = col_widths[i]
        set_cell_background(hdr_cells[i], "1B4D3E")
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = p.add_run(title)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        run.font.size = Pt(10)
        set_cell_margins(hdr_cells[i], top=100, bottom=100, left=100, right=100)

    rows_data = [
        ("Diamond / Gem Scavenger Hunt", "Site Activity & Sales", "Low - Med", "⭐⭐⭐⭐ (High)"),
        ("1-Tap Social Unlock / Share", "Follower Growth & Reach", "Very Low", "⭐⭐⭐ (Medium)"),
        ("Spin-to-Win Mystery Wheel", "Lead Generation & Sales", "Medium", "⭐⭐⭐⭐ (High)"),
        ("Rate My Plate (TikTok Roast)", "Site Traffic & Consults", "Low", "⭐⭐⭐⭐⭐ (Explosive)"),
        ("Golden Ticket Meal Plan", "Direct Plan Sales", "Low", "⭐⭐⭐⭐⭐ (Explosive)"),
        ("TikTok Live Flash Consult Blitz", "Immediate Consult Bookings", "Very Low", "⭐⭐⭐⭐⭐ (Explosive)"),
        ("Community Goal Unlock Bar", "Viral Traffic & Buzz", "Medium", "⭐⭐⭐ (Medium)")
    ]

    for row_idx, data in enumerate(rows_data, start=1):
        row_cells = table.rows[row_idx].cells
        bg_color = "F9FAFB" if row_idx % 2 == 0 else "FFFFFF"
        for i, val in enumerate(data):
            row_cells[i].width = col_widths[i]
            set_cell_background(row_cells[i], bg_color)
            p = row_cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run(val)
            run.font.size = Pt(9.5)
            if i == 0:
                run.font.bold = True
                run.font.color.rgb = COLOR_PRIMARY
            set_cell_margins(row_cells[i], top=80, bottom=80, left=100, right=100)

    p_space = doc.add_paragraph()
    p_space.paragraph_format.space_before = Pt(10)

    add_divider()

    add_h1("5. Recommended Action Plan & Next Steps")
    add_callout(
        "💡 Recommended Campaign Combination for Maximum Impact",
        "We recommend combining 2 complementary ideas for the ultimate launch: (1) The Diamond Scavenger Hunt on dietwithdee.org to drive immediate website traffic, paired with (2) The TikTok Live Flash Consultation Blitz to turn that traffic directly into paid 1-on-1 consultations."
    )

    p = doc.add_paragraph()
    p.add_run("Next Steps for Implementation:\n").bold = True
    p.add_run("1. Dee Reviews & Approves Selected Idea Combination.\n")
    p.add_run("2. Web Update: Add Diamond icons or landing page banner on dietwithdee.org.\n")
    p.add_run("3. Social Announcement: Post a 30-second TikTok teaser announcing the celebration campaign and giveaway rules.\n")
    p.add_run("4. Execution & Live Event: Run campaign for 5 days culminating in a celebratory TikTok Live stream.")

    # Save document
    doc.save(output_path)
    print(f"Successfully generated proposal document at: {output_path}")

if __name__ == "__main__":
    create_proposal_doc()
